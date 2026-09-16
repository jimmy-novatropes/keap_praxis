#!/usr/bin/env python3
"""
Export every Keap opportunity (deal) with the tags of its linked contact.

Like companies, opportunities don't carry tags themselves in Keap --
only Contacts do. So each opportunity's "tags" here are just its
linked contact's tags.

Setup:
    pip install requests openpyxl python-dotenv
    Uses the same .env as "keap tags.py":
        KEAP_API_KEY=the key you generated

Usage:
    python keap_opportunity_tags.py [--output keap_opportunity_tags_export.xlsx]

Auth: Keap REST API v1, key-based auth via the `X-Keap-API-Key` header.
Docs: https://developer.infusionsoft.com/docs/rest/

NOTE on unverified assumptions (same caveat style as "get subs.py" --
matches the v1 docs but hasn't been re-confirmed against a live
response as of this writing):
    - GET /opportunities paginates with limit/offset, list key
      "opportunities", same as /contacts.
    - Each opportunity has a nested "contact" object with at least an
      "id" (e.g. {"id": 123}), and a nested "stage" object with "id"
      and "label"/"name". This script tries a couple of key names for
      the stage label; if it prints "[unknown stage field]" for every
      row, run with --debug once to print a raw record and fix
      STAGE_NAME_KEYS below.
"""

import os
import re
import html
import sys
import time
import argparse
import requests
from openpyxl import Workbook

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    sys.exit("Missing dependency: run `pip install python-dotenv` (or `pip install -r requirements.txt`).")

BASE_URL = "https://api.infusionsoft.com/crm/rest/v1"
PAGE_SIZE = 200
STAGE_NAME_KEYS = ("label", "name")  # tried in order against the "stage" object


def get_api_key():
    key = os.environ.get("KEAP_API_KEY")
    if not key:
        sys.exit("Set KEAP_API_KEY in a .env file (or as an env var) before running this.")
    return key


def keap_get(session, path, params=None, max_retries=5):
    """GET against the Keap API with basic 429/5xx retry + backoff."""
    url = f"{BASE_URL}{path}" if path.startswith("/") else path
    resp = None
    for attempt in range(max_retries):
        resp = session.get(url, params=params, timeout=30)
        if resp.status_code == 429:
            wait = int(resp.headers.get("Retry-After", 5) or 5)
            print(f"  Rate limited, waiting {wait}s...")
            time.sleep(wait)
            continue
        if resp.status_code >= 500:
            wait = 2 ** attempt
            print(f"  Server error {resp.status_code}, retrying in {wait}s...")
            time.sleep(wait)
            continue
        resp.raise_for_status()
        return resp.json()
    resp.raise_for_status()


def fetch_all_tags(session):
    """Fetch every tag definition (id -> name)."""
    tags = {}
    offset = 0
    while True:
        data = keap_get(session, "/tags", params={"limit": 1000, "offset": offset})
        batch = data.get("tags", [])
        for t in batch:
            tags[t["id"]] = html.unescape(t.get("name", "") or "")
        offset += len(batch)
        if len(batch) < 1000:
            break
    print(f"Fetched {len(tags)} tag definitions.")
    return tags


def slugify_tag(name):
    """Best-effort HubSpot internal-name suggestion (lowercase, underscores).
    HubSpot auto-generates its own internal value when you paste labels in,
    so treat this as a preview/sanity-check rather than the source of truth."""
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    return slug or "tag"


def write_tag_reference(tag_names, wb, txt_path):
    """
    Write every distinct Keap tag name to:
      1. A plain .txt file, one tag per line, sorted -- paste this directly
         into HubSpot's "Add options" box when creating a multi-checkbox or
         dropdown select property to pre-populate every option at once.
      2. An "All Tags" sheet in the same workbook, with a suggested
         HubSpot internal value alongside each label, for reference.
    Returns the sorted list of unique tag names.
    """
    unique_names = sorted({n for n in tag_names.values() if n}, key=str.lower)

    with open(txt_path, "w", encoding="utf-8") as f:
        for name in unique_names:
            f.write(name + "\n")
    print(f"Wrote {len(unique_names)} unique tag names to {txt_path}")

    ws = wb.create_sheet("All Tags")
    ws.append(["tag_name (HubSpot label)", "suggested_internal_value"])
    seen_slugs = {}
    for name in unique_names:
        slug = slugify_tag(name)
        if slug in seen_slugs:
            seen_slugs[slug] += 1
            slug = f"{slug}_{seen_slugs[slug]}"
        else:
            seen_slugs[slug] = 0
        ws.append([name, slug])

    return unique_names


def fetch_all_contacts_indexed(session):
    """Fetch every contact once; return {contact_id: {email, first, last, tag_ids}}."""
    index = {}
    offset = 0
    while True:
        params = {"limit": PAGE_SIZE, "offset": offset, "optional_properties": "tag_ids"}
        data = keap_get(session, "/contacts", params=params)
        batch = data.get("contacts", [])
        for c in batch:
            index[c["id"]] = {
                "email": primary_email(c),
                "first": c.get("given_name", "") or "",
                "last": c.get("family_name", "") or "",
                "tag_ids": c.get("tag_ids") or [],
            }
        print(f"  Indexed {len(index)} contacts so far...")
        if len(batch) < PAGE_SIZE:
            break
        offset += PAGE_SIZE
    return index


def primary_email(contact):
    emails = contact.get("email_addresses") or []
    for e in emails:
        if e.get("field") == "EMAIL1":
            return e.get("email", "")
    return emails[0]["email"] if emails else ""


def fetch_all_opportunities(session):
    opportunities = []
    offset = 0
    while True:
        data = keap_get(session, "/opportunities", params={"limit": PAGE_SIZE, "offset": offset})
        batch = data.get("opportunities", [])
        opportunities.extend(batch)
        print(f"  Fetched {len(opportunities)} opportunities so far...")
        if len(batch) < PAGE_SIZE:
            break
        offset += PAGE_SIZE
    return opportunities


def stage_label(opp):
    stage = opp.get("stage") or {}
    for key in STAGE_NAME_KEYS:
        if stage.get(key):
            return stage[key]
    return "[unknown stage field]" if stage else ""


def main():
    parser = argparse.ArgumentParser(description="Export Keap opportunities + their contacts' tags to Excel.")
    parser.add_argument("--output", default="keap_opportunity_tags_export.xlsx")
    parser.add_argument("--tags-txt", default="keap_all_tags_for_hubspot.txt",
                         help="Plain-text file listing every distinct tag name, one per line.")
    parser.add_argument("--debug", action="store_true", help="Print one raw opportunity record and exit.")
    args = parser.parse_args()

    api_key = get_api_key()
    session = requests.Session()
    session.headers.update({"X-Keap-API-Key": api_key, "Accept": "application/json"})

    if args.debug:
        data = keap_get(session, "/opportunities", params={"limit": 1, "offset": 0})
        import json
        print(json.dumps(data.get("opportunities", [None])[0], indent=2))
        return

    print("Fetching tag definitions...")
    tag_names = fetch_all_tags(session)

    print("Indexing contacts (tags + basic info)...")
    contacts = fetch_all_contacts_indexed(session)
    print(f"Indexed {len(contacts)} contacts total.")

    print("Fetching opportunities...")
    opportunities = fetch_all_opportunities(session)
    print(f"Fetched {len(opportunities)} opportunities total.")

    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Opportunities"
    ws1.append(["opportunity_id", "opportunity_title", "stage", "contact_id",
                "email", "first_name", "last_name", "tags"])

    write_tag_reference(tag_names, wb, args.tags_txt)

    ws2 = wb.create_sheet("Opportunity-Tag Pairs")
    ws2.append(["opportunity_id", "opportunity_title", "tag_id", "tag_name"])

    unmatched_tag_ids = set()
    no_contact = 0

    for opp in opportunities:
        oid = opp.get("id")
        title = opp.get("opportunity_title") or opp.get("title") or ""
        contact_id = (opp.get("contact") or {}).get("id")
        contact = contacts.get(contact_id, {})
        if not contact_id:
            no_contact += 1

        names = []
        for tid in contact.get("tag_ids", []):
            name = tag_names.get(tid)
            if name is None:
                unmatched_tag_ids.add(tid)
                name = f"[unknown tag {tid}]"
            names.append(name)
            ws2.append([oid, title, tid, name])

        ws1.append([
            oid, title, stage_label(opp), contact_id,
            contact.get("email", ""), contact.get("first", ""), contact.get("last", ""),
            "; ".join(sorted(names)),
        ])

    wb.save(args.output)
    print(f"\nSaved {len(opportunities)} opportunities to {args.output}")
    if no_contact:
        print(f"Note: {no_contact} opportunity(ies) had no linked contact.")
    if unmatched_tag_ids:
        print(f"Warning: {len(unmatched_tag_ids)} tag id(s) had no matching name: {sorted(unmatched_tag_ids)}")


if __name__ == "__main__":
    main()
